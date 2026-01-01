import hashlib
import os
import re
import tkinter as tk
from tkinter import filedialog, messagebox
from tkinter import ttk

from docx import Document


MODE_PRESETS = {
    "Conservative": {
        "replace_contacts": True,
        "replace_persons": True,
        "replace_brands": False,
        "replace_geo": False,
        "replace_numbers": False,
    },
    "Balanced": {
        "replace_contacts": True,
        "replace_persons": True,
        "replace_brands": True,
        "replace_geo": True,
        "replace_numbers": True,
    },
    "Strict": {
        "replace_contacts": True,
        "replace_persons": True,
        "replace_brands": True,
        "replace_geo": True,
        "replace_numbers": True,
    },
}


BRAND_TERMS = [
    "Google",
    "Microsoft",
    "Apple",
    "Amazon",
    "Meta",
    "Netflix",
    "Tesla",
]

BRAND_TERMS_STRICT = BRAND_TERMS + [
    "OpenAI",
    "Salesforce",
    "Adobe",
    "Oracle",
]

GEO_TERMS = [
    "New York",
    "London",
    "Paris",
    "Berlin",
    "Tokyo",
    "USA",
    "United States",
    "Germany",
]

GEO_TERMS_STRICT = GEO_TERMS + [
    "Canada",
    "Spain",
    "Italy",
    "Australia",
    "India",
]

STOP_WORDS = {
    "brief",
    "transcript",
    "meeting",
    "notes",
    "note",
    "interview",
    "document",
    "project",
    "report",
    "nda",
    "draft",
    "final",
    "file",
    "version",
    "summary",
    "analysis",
    "output",
}


def make_short_hash(text):
    return hashlib.sha1(text.encode("utf-8")).hexdigest()[:6]


def sanitize_windows_name(text):
    sanitized = re.sub(r'[<>:"/\\\\|?*]', "_", text)
    sanitized = re.sub(r"[\x00-\x1f]", "", sanitized)
    sanitized = sanitized.strip().strip(".")
    return sanitized


def make_safe_project_id(selected_files, manual_label):
    label = sanitize_windows_name(manual_label.strip()) if manual_label else ""
    if label:
        return label[:50]
    hash_source = "|".join(selected_files) if selected_files else "empty"
    return f"Project_{make_short_hash(hash_source)}"


def extract_sensitive_tokens(paths):
    tokens = set()
    for path in paths:
        stem = os.path.splitext(os.path.basename(path))[0]
        for part in re.split(r"[^A-Za-z0-9]+", stem):
            part_clean = part.strip()
            if len(part_clean) < 3:
                continue
            if part_clean.lower() in STOP_WORDS:
                continue
            tokens.add(part_clean)
    return sorted(tokens, key=lambda value: (-len(value), value.lower()))


def mask_sensitive_names(text, known_tokens):
    if not text:
        return text

    token_map = {}
    project_counter = 1
    brand_counter = 1

    def replace_known(match):
        nonlocal project_counter
        original = match.group(0)
        key = original.lower()
        token = token_map.get(key)
        if token is None:
            token = f"PROJECT_{project_counter}"
            project_counter += 1
            token_map[key] = token
        return token

    for token in known_tokens:
        if not token:
            continue
        pattern = rf"\b{re.escape(token)}\b"
        text = re.sub(pattern, replace_known, text, flags=re.IGNORECASE)

    def replace_all_caps(match):
        nonlocal brand_counter
        original = match.group(0)
        token = token_map.get(original)
        if token is None:
            token = f"BRAND_{brand_counter}"
            brand_counter += 1
            token_map[original] = token
        return token

    text = re.sub(r"\b[A-Z]{3,}\b", replace_all_caps, text)
    return text


def sanitize_text(
    text,
    replace_contacts,
    replace_persons,
    replace_brands,
    replace_geo,
    replace_numbers,
    strict=False,
):
    redaction_map = []
    domain_replacements = 0
    brand_replacements = 0
    domain_token_by_base = {}
    domain_token_by_domain = {}
    brand_token_by_term = {}
    domain_token_counter = 1
    brand_token_counter = 1

    def apply_replacement(pattern, replacement, label):
        nonlocal text

        def _replace(match):
            original = match.group(0)
            redaction_map.append((original, replacement, label))
            return replacement

        text = re.sub(pattern, _replace, text, flags=re.IGNORECASE | re.MULTILINE)

    def apply_replacement_func(pattern, replace_func, label):
        nonlocal text

        def _replace(match):
            original = match.group(0)
            replacement = replace_func(match)
            redaction_map.append((original, replacement, label))
            return replacement

        text = re.sub(pattern, _replace, text, flags=re.IGNORECASE | re.MULTILINE)

    if replace_contacts:
        email_pattern = r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}"
        phone_pattern = r"\+?\d[\d\s().-]{7,}\d"
        apply_replacement(email_pattern, "[CONTACT]", "contact_email")
        apply_replacement(phone_pattern, "[CONTACT]", "contact_phone")

    if replace_numbers:
        apply_replacement(r"\b\d{2,}\b", "[NUMBER]", "number")

    if replace_persons:
        name_pattern = r"\b([A-Z][a-z]+\s+[A-Z][a-z]+)\b"
        apply_replacement(name_pattern, "[PERSON]", "person")

    if replace_brands:
        domain_pattern = r"\b(?:[a-z0-9-]+\.)+[a-z]{2,}\b"
        for match in re.finditer(domain_pattern, text, flags=re.IGNORECASE | re.MULTILINE):
            domain = match.group(0)
            domain_lower = domain.lower()
            labels = domain_lower.split(".")
            base = labels[-2] if len(labels) >= 2 else labels[0]
            if base not in domain_token_by_base:
                domain_token_by_base[base] = f"DOMAIN_{domain_token_counter}"
                domain_token_counter += 1
            domain_token_by_domain[domain_lower] = domain_token_by_base[base]

        def domain_replace(match):
            nonlocal domain_replacements, domain_token_counter
            domain = match.group(0)
            token = domain_token_by_domain.get(domain.lower())
            if token is None:
                labels = domain.lower().split(".")
                base = labels[-2] if len(labels) >= 2 else labels[0]
                token = domain_token_by_base.get(base)
                if token is None:
                    token = f"DOMAIN_{domain_token_counter}"
                    domain_token_by_base[base] = token
                    domain_token_counter += 1
                domain_token_by_domain[domain.lower()] = token
            domain_replacements += 1
            return token

        apply_replacement_func(domain_pattern, domain_replace, "brand_domain")

        terms = BRAND_TERMS_STRICT if strict else BRAND_TERMS
        for term in terms:
            def brand_replace(match, term_key=term.lower()):
                nonlocal brand_replacements, brand_token_counter
                token = domain_token_by_base.get(term_key)
                if token is None:
                    token = brand_token_by_term.get(term_key)
                    if token is None:
                        token = f"BRAND_{brand_token_counter}"
                        brand_token_counter += 1
                        brand_token_by_term[term_key] = token
                brand_replacements += 1
                return token

            apply_replacement_func(rf"\b{re.escape(term)}\b", brand_replace, "brand")

        for base, token in domain_token_by_base.items():
            def base_replace(match, token=token):
                nonlocal brand_replacements
                brand_replacements += 1
                return token

            apply_replacement_func(rf"\b{re.escape(base)}\b", base_replace, "brand")

    if replace_geo:
        terms = GEO_TERMS_STRICT if strict else GEO_TERMS
        for term in terms:
            apply_replacement(rf"\b{re.escape(term)}\b", "[GEO]", "geo")

    qa_notes = []
    if not text.strip():
        qa_notes.append("Sanitized text is empty after processing.")

    return text, redaction_map, qa_notes, domain_replacements, brand_replacements


class BriefLegitimizerApp:
    def __init__(self, root):
        self.root = root
        self.root.title("NDA Firewall")

        self.file_path = tk.StringVar()
        self.safe_mode_var = tk.BooleanVar(value=True)
        self.project_label_var = tk.StringVar()
        self.mode_var = tk.StringVar(value="Balanced")

        self.replace_contacts_var = tk.BooleanVar()
        self.replace_persons_var = tk.BooleanVar()
        self.replace_brands_var = tk.BooleanVar()
        self.replace_geo_var = tk.BooleanVar()
        self.replace_numbers_var = tk.BooleanVar()

        self.selected_file = ""
        self.original_text = ""
        self.sanitized_text = ""
        self.redaction_map = []
        self.qa_notes = []

        self._build_ui()
        self.apply_mode_policy()
        self.mode_var.trace_add("write", self._on_mode_change)
        self.safe_mode_var.trace_add("write", self._on_safe_mode_change)
        self._update_unsafe_warning()

    def _build_ui(self):
        top_frame = ttk.Frame(self.root, padding=10)
        top_frame.pack(fill="x")

        file_label = ttk.Label(top_frame, text="Input")
        file_label.grid(row=0, column=0, sticky="w")

        file_entry = ttk.Entry(top_frame, textvariable=self.file_path, width=50)
        file_entry.grid(row=0, column=1, padx=5, sticky="we")

        browse_button = ttk.Button(top_frame, text="Browse", command=self.browse_file)
        browse_button.grid(row=0, column=2, padx=5)

        preview_button = ttk.Button(top_frame, text="Preview", command=self.preview_file)
        preview_button.grid(row=0, column=3, padx=5)

        sanitize_button = ttk.Button(
            top_frame, text="Legitimize", command=self.run_sanitization
        )
        sanitize_button.grid(row=0, column=4, padx=5)

        export_button = ttk.Button(top_frame, text="Save SAFE TXT...", command=self.export_txt)
        export_button.grid(row=0, column=5, padx=5)

        mode_frame = ttk.Frame(top_frame)
        mode_frame.grid(row=1, column=0, columnspan=6, sticky="w", pady=(8, 0))

        mode_label = ttk.Label(mode_frame, text="Mode:")
        mode_label.pack(side="left")

        mode_options = [
            ("Conservative", "Conservative"),
            ("Balanced (default)", "Balanced"),
            ("Strict", "Strict"),
        ]
        for label, value in mode_options:
            ttk.Radiobutton(
                mode_frame, text=label, value=value, variable=self.mode_var
            ).pack(side="left", padx=(6, 0))

        safety_frame = ttk.Frame(top_frame)
        safety_frame.grid(row=2, column=0, columnspan=6, sticky="w", pady=(8, 0))

        ttk.Checkbutton(
            safety_frame,
            text="SAFE mode (hide file/project names)",
            variable=self.safe_mode_var,
        ).pack(side="left")

        self.unsafe_warning_label = ttk.Label(
            safety_frame,
            text="UNSAFE MODE: may reveal client/brand",
            foreground="red",
        )
        self.unsafe_warning_label.pack(side="left", padx=(12, 0))

        project_frame = ttk.Frame(top_frame)
        project_frame.grid(row=3, column=0, columnspan=6, sticky="we", pady=(6, 0))

        ttk.Label(project_frame, text="Project label (optional):").pack(side="left")
        ttk.Entry(project_frame, textvariable=self.project_label_var, width=40).pack(
            side="left", padx=(6, 0), fill="x", expand=True
        )

        top_frame.columnconfigure(1, weight=1)

        policy_frame = ttk.LabelFrame(self.root, text="NDA Firewall Policy (Advanced)")
        policy_frame.pack(fill="x", padx=10, pady=(0, 10))

        ttk.Label(
            policy_frame,
            text="Mode sets the preset. Checkboxes can be adjusted manually (Advanced).",
        ).pack(anchor="w", padx=8, pady=(4, 0))

        checks_frame = ttk.Frame(policy_frame)
        checks_frame.pack(fill="x", padx=8, pady=6)

        ttk.Checkbutton(
            checks_frame,
            text="Replace contacts",
            variable=self.replace_contacts_var,
        ).grid(row=0, column=0, sticky="w")
        ttk.Checkbutton(
            checks_frame,
            text="Replace persons",
            variable=self.replace_persons_var,
        ).grid(row=0, column=1, sticky="w", padx=10)
        ttk.Checkbutton(
            checks_frame,
            text="Replace brands",
            variable=self.replace_brands_var,
        ).grid(row=1, column=0, sticky="w", pady=4)
        ttk.Checkbutton(
            checks_frame,
            text="Replace geo",
            variable=self.replace_geo_var,
        ).grid(row=1, column=1, sticky="w", padx=10)
        ttk.Checkbutton(
            checks_frame,
            text="Replace numbers",
            variable=self.replace_numbers_var,
        ).grid(row=2, column=0, sticky="w")

        text_frame = ttk.Frame(self.root, padding=10)
        text_frame.pack(fill="both", expand=True)

        self.preview_text = tk.Text(text_frame, height=10, wrap="word")
        self.preview_text.pack(side="left", fill="both", expand=True, padx=(0, 5))

        self.sanitized_text_widget = tk.Text(text_frame, height=10, wrap="word")
        self.sanitized_text_widget.pack(side="left", fill="both", expand=True, padx=(5, 0))

        log_frame = ttk.Frame(self.root, padding=10)
        log_frame.pack(fill="both")

        ttk.Label(log_frame, text="Log").pack(anchor="w")
        self.log_text = tk.Text(log_frame, height=6, wrap="word")
        self.log_text.pack(fill="both", expand=True)

    def _on_mode_change(self, *_):
        self.apply_mode_policy()

    def _on_safe_mode_change(self, *_):
        self._refresh_file_display()
        self._update_unsafe_warning()

    def apply_mode_policy(self):
        preset = MODE_PRESETS.get(self.mode_var.get(), MODE_PRESETS["Balanced"])
        self.replace_contacts_var.set(preset["replace_contacts"])
        self.replace_persons_var.set(preset["replace_persons"])
        self.replace_brands_var.set(preset["replace_brands"])
        self.replace_geo_var.set(preset["replace_geo"])
        self.replace_numbers_var.set(preset["replace_numbers"])

    def _update_unsafe_warning(self):
        if self.safe_mode_var.get():
            self.unsafe_warning_label.configure(text="")
        else:
            self.unsafe_warning_label.configure(text="UNSAFE MODE: may reveal client/brand")

    def _refresh_file_display(self):
        if not self.selected_file:
            self.file_path.set("")
            return
        if self.safe_mode_var.get():
            self.file_path.set("File #1")
        else:
            self.file_path.set(self.selected_file)

    def _current_file_label(self):
        return "File #1" if self.safe_mode_var.get() else self.selected_file

    def _get_project_label(self):
        return make_safe_project_id([self.selected_file], self.project_label_var.get())

    def _get_known_tokens(self):
        tokens = extract_sensitive_tokens([self.selected_file])
        manual_label = self.project_label_var.get().strip()
        if manual_label:
            manual_tokens = extract_sensitive_tokens([manual_label])
            tokens.extend(manual_tokens)
        return sorted(set(tokens), key=lambda value: (-len(value), value.lower()))

    def browse_file(self):
        path = filedialog.askopenfilename(
            title="Select brief file",
            filetypes=[
                ("All files", "*.*"),
                ("Documents", "*.docx *.txt"),
                ("Word Document", "*.docx"),
                ("Text", "*.txt"),
            ],
        )
        if path:
            self.selected_file = path
            self._refresh_file_display()
            self._update_unsafe_warning()

    def _load_text_from_file(self, path):
        if path.lower().endswith(".docx"):
            document = Document(path)
            return "\n".join(paragraph.text for paragraph in document.paragraphs)
        with open(path, "r", encoding="utf-8") as handle:
            return handle.read()

    def preview_file(self):
        path = self.selected_file
        if not path:
            messagebox.showwarning("Missing file", "Select a file first.")
            return
        if not os.path.exists(path):
            messagebox.showerror("Missing file", "File not found.")
            return
        self.original_text = self._load_text_from_file(path)
        self.preview_text.delete("1.0", tk.END)
        if self.safe_mode_var.get():
            masked_preview = mask_sensitive_names(
                self.original_text, self._get_known_tokens()
            )
            self.preview_text.insert(tk.END, masked_preview)
            self._log(f"Preview loaded for {self._current_file_label()}.")
        else:
            self.preview_text.insert(tk.END, self.original_text)
            self._log(f"Preview loaded: {path}")

    def _log(self, message):
        self.log_text.insert(tk.END, message + "\n")
        self.log_text.see(tk.END)

    def _collect_policy(self):
        return {
            "replace_contacts": self.replace_contacts_var.get(),
            "replace_persons": self.replace_persons_var.get(),
            "replace_brands": self.replace_brands_var.get(),
            "replace_geo": self.replace_geo_var.get(),
            "replace_numbers": self.replace_numbers_var.get(),
        }

    def run_sanitization(self):
        if not self.original_text.strip():
            self.preview_file()
        if not self.original_text.strip():
            return

        policy = self._collect_policy()
        mode = self.mode_var.get()
        strict = mode == "Strict"

        policy_log = (
            f"[POLICY] mode={mode} | brands={policy['replace_brands']} "
            f"persons={policy['replace_persons']} geo={policy['replace_geo']} "
            f"numbers={policy['replace_numbers']} contacts={policy['replace_contacts']}"
        )
        self._log(policy_log)

        (
            self.sanitized_text,
            self.redaction_map,
            self.qa_notes,
            domain_replacements,
            brand_replacements,
        ) = sanitize_text(
            self.original_text,
            strict=strict,
            **policy,
        )

        self.sanitized_text_widget.delete("1.0", tk.END)
        if self.safe_mode_var.get():
            self.sanitized_text = mask_sensitive_names(
                self.sanitized_text, self._get_known_tokens()
            )
        self.sanitized_text_widget.insert(tk.END, self.sanitized_text)
        self._log(
            f"[BRAND_REDACTION] domains={domain_replacements} brands={brand_replacements}"
        )
        self._log("NDA Firewall complete.")

    def export_txt(self):
        if not self.sanitized_text:
            messagebox.showwarning("Missing output", "Run legitimization first.")
            return

        project_label = self._get_project_label()
        safe_mode = self.safe_mode_var.get()
        if safe_mode:
            initial_name = f"sanitized_output__{project_label}.txt"
        else:
            stem = os.path.splitext(os.path.basename(self.selected_file or "output"))[0]
            initial_name = f"{stem}_sanitized.txt"

        path = filedialog.asksaveasfilename(
            title="Save SAFE TXT",
            defaultextension=".txt",
            initialfile=initial_name,
            filetypes=[("All files", "*.*"), ("Text File", "*.txt")],
        )
        if not path:
            return

        header = "NDA Firewall Output (SAFE)" if safe_mode else "NDA Firewall Output (UNSAFE)"
        if project_label:
            header = f"{header}\nProject: {project_label}"
        content = f"{header}\n\n{self.sanitized_text}"
        with open(path, "w", encoding="utf-8") as handle:
            handle.write(content)
        if safe_mode:
            self._log(f"SAFE TXT exported: {os.path.basename(path)}")
        else:
            self._log(f"SAFE TXT exported: {path}")


if __name__ == "__main__":
    root = tk.Tk()
    app = BriefLegitimizerApp(root)
    root.mainloop()
